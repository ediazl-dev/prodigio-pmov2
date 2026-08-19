#!/usr/bin/env python3
"""Safe Deal-ID UPSERT from an exported financial XLSX into the target PMO database."""
import argparse
import json
import os
import re
import sys
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation

import openpyxl
import pymysql

SHEET_NAME = "Artefactos_proyectos"
FIELD_MAP = {
    "Estado Proyecto": "estadoProyecto", "Proyecto": "projectName", "Cliente": "clientName", "PM": "pm",
    "valor_venta_uf": "valorVentaUF", "Presupuesto_UF": "presupuestoUF", "Utilizado_UF": "utilizadoUF",
    "Utilizado_UF_porc": "utilizadoUFPorc", "Presupuesto_HH": "presupuestoHH", "Capacity_HH": "capacityHH",
    "HH_porc_utilizado": "hhPorcUtilizado", "Margen_bruto_nota_venta (UF)": "margenBrutoNotaVentaUF",
    "Porcentaje_avance_proyecto": "porcentajeAvanceProyecto", "Costo_Proyectado_UF_(segun avance)": "costoProyectadoUF",
    "Margen_Proyectado UF_(segun avance)": "margenProyectadoUF", "Margen_Proyectado_% (segun avance)": "margenProyectadoPorc",
    "Margen_Target_porc": "margenTargetPorc", "Capacity U": "capacityU", "Planificado_UF": "planificadoUF",
    "Proyectado_UF": "proyectadoUF", "margen_proyectado_segun_capacity": "margenProyectadoSegunCapacity",
    "Notas": "notas", "Otros costos (UF)": "otrosCostosUF", "Linea_negocio": "lineaNegocio",
}
NUMERIC_FIELDS = {value for key, value in FIELD_MAP.items() if key not in {"Estado Proyecto", "Proyecto", "Cliente", "PM", "Notas", "Linea_negocio"}}

def fail(message):
    raise RuntimeError(message)

def database_config():
    url = os.environ.get("DATABASE_URL", "")
    if not url:
        fail("DATABASE_URL no está configurada en el entorno destino")
    cleaned = re.sub(r"[?&](ssl|sslaccept)=[^&]*", "", url)
    match = re.match(r"mysql(?:2)?://([^:]+):([^@]+)@([^:/]+):?(\d+)?/([^?]+)", cleaned)
    if not match:
        fail("DATABASE_URL tiene un formato no compatible")
    user, password, host, port, database = match.groups()
    return {"host": host, "port": int(port or 4000), "user": user, "password": password, "database": database.split("?")[0]}

def text(value):
    if value is None:
        return None
    value = str(value).strip()
    return None if value in {"", "None", "N/A"} or value.startswith("#") else value

def number(value):
    if value is None or str(value).strip() in {"", "None", "N/A"}:
        return None
    try:
        return Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        return None

def parse_workbook(path):
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        fail(f"No existe la hoja requerida: {SHEET_NAME}")
    sheet = workbook[SHEET_NAME]
    headers = [text(cell.value) or "" for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    index = {header: position for position, header in enumerate(headers) if header}
    if "Deal" not in index:
        fail("La hoja no contiene la columna obligatoria Deal")
    missing = [header for header in FIELD_MAP if header not in index]
    if missing:
        fail("Faltan columnas financieras requeridas: " + ", ".join(missing))
    rows, duplicate_deals = [], set()
    seen = set()
    for cells in sheet.iter_rows(min_row=2, values_only=True):
        deal_id = text(cells[index["Deal"]])
        if not deal_id:
            continue
        if not deal_id.startswith("Deal"):
            continue
        if deal_id in seen:
            duplicate_deals.add(deal_id)
            continue
        seen.add(deal_id)
        record = {"dealId": deal_id}
        for source, target in FIELD_MAP.items():
            raw = cells[index[source]]
            record[target] = number(raw) if target in NUMERIC_FIELDS else text(raw)
        rows.append(record)
    if not rows:
        fail("No se encontraron Deal IDs válidos; no se escribirá nada")
    if duplicate_deals:
        fail("Deal IDs duplicados en la planilla: " + ", ".join(sorted(duplicate_deals)))
    return rows

def verify_schema(cursor):
    cursor.execute("SHOW COLUMNS FROM financial_data")
    available = {row[0] for row in cursor.fetchall()}
    required = {"dealId", *FIELD_MAP.values(), "syncedAt"}
    missing = sorted(required - available)
    if missing:
        fail("La base destino no tiene la migración financiera requerida: " + ", ".join(missing))

def run(rows, apply):
    config = database_config()
    connection = pymysql.connect(**config, ssl={"ssl": {}}, charset="utf8mb4", autocommit=False)
    fields = list(FIELD_MAP.values())
    outcome = {"timestampUtc": datetime.now(timezone.utc).isoformat(), "inputDeals": len(rows), "insert": 0, "update": 0, "status": "dry_run"}
    try:
        with connection.cursor() as cursor:
            verify_schema(cursor)
            existing = set()
            deal_ids = [row["dealId"] for row in rows]
            for offset in range(0, len(deal_ids), 500):
                batch = deal_ids[offset:offset + 500]
                cursor.execute("SELECT dealId FROM financial_data WHERE dealId IN (" + ",".join(["%s"] * len(batch)) + ")", batch)
                existing.update(row[0] for row in cursor.fetchall())
            outcome["insert"] = sum(1 for deal in deal_ids if deal not in existing)
            outcome["update"] = len(deal_ids) - outcome["insert"]
            if not apply:
                return outcome
            columns = ["dealId", *fields]
            insert_columns = ", ".join(f"`{column}`" for column in columns)
            placeholders = ", ".join(["%s"] * len(columns))
            update_clause = ", ".join(f"`{field}`=VALUES(`{field}`)" for field in fields) + ", `syncedAt`=CURRENT_TIMESTAMP"
            sql = f"INSERT INTO financial_data ({insert_columns}) VALUES ({placeholders}) ON DUPLICATE KEY UPDATE {update_clause}"
            for row in rows:
                cursor.execute(sql, [row.get(column) for column in columns])
            connection.commit()
            outcome["status"] = "applied"
            return outcome
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()

def main():
    parser = argparse.ArgumentParser(description="Sincroniza datos financieros hacia una instancia PMO destino")
    parser.add_argument("--sheet", required=True, help="Ruta XLSX previamente exportada desde Google Drive")
    parser.add_argument("--apply", action="store_true", help="Aplica el UPSERT; sin esta opción sólo ejecuta dry-run")
    args = parser.parse_args()
    if not os.path.isfile(args.sheet):
        fail("No existe el XLSX indicado")
    result = run(parse_workbook(args.sheet), args.apply)
    print(json.dumps(result, ensure_ascii=False, default=str))

if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(json.dumps({"status": "error", "error": str(error)}, ensure_ascii=False), file=sys.stderr)
        sys.exit(1)
